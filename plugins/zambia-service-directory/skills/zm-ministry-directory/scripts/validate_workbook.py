#!/usr/bin/env python3
"""Check a ministry workbook against the fixed format and its own totals.

Usage: validate_workbook.py <workbook.xlsx> [--spec templates/workbook_spec.json] [--out report.json]

Exit 0 = no errors, 1 = errors, 2 = cannot open.

Rule ids:
  W-SHEETS    sheet names/order
  W-HEADER    header text (Overview table, Agencies, every Service Passport section)
  W-SECTION   section pattern: title -> header -> body -> spacer; title merged A:K
  W-MERGE     merged range that is not a section title row
  W-STYLE     font/fill/border/alignment differs from the spec role (warn per column, error on header/title)
  W-LAYOUT    widths / freeze panes / gridlines differ from spec (warn)
  W-TOTAL     Overview totals or per-agency counts differ from actual rows; Notes count mismatch
  W-EMPTY     empty required cell
  W-TERM      deprecated placeholder (e.g. 'Not specified')
  W-DUP       duplicate service rows (same agency+name, or same eServices link twice)
  W-FORMULA   formula error value (#REF!, #VALUE!, ...)
  W-CONSIST   ministry / agency names disagree between sheets
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

import zmcontract as c
from zmlog import dump_json, get_logger, load_json, templates_dir
from zmstyle import style_diff, style_of
from zmworkbook import SECTION_PREFIX, eservices_id_from_link, read_workbook

log = get_logger("validate_workbook")
ERROR_VALUES = ("#REF!", "#VALUE!", "#NAME?", "#DIV/0!", "#N/A", "#NUM!", "#NULL!")
STYLE_KEYS = ("font", "fill", "border", "alignment")


class Report:
    def __init__(self):
        self.items: list[dict] = []

    def add(self, rule, severity, sheet, cell, message):
        self.items.append({"rule": rule, "severity": severity, "sheet": sheet, "cell": cell, "message": message})
        log.debug("%s %s %s!%s %s", rule, severity, sheet, cell, message)


def _style_check(rep: Report, ws, row: int, expected: list[dict], severity: str, role: str) -> None:
    for col, exp in enumerate(expected, start=1):
        act = style_of(ws.cell(row, col))
        diffs = style_diff({k: act[k] for k in STYLE_KEYS}, {k: exp[k] for k in STYLE_KEYS})
        if diffs:
            rep.add("W-STYLE", severity, ws.title, f"{get_column_letter(col)}{row}",
                    f"{role} style differs: {'; '.join(diffs[:3])}")


def check_layout(rep: Report, ws, sheet_spec: dict) -> None:
    if ws.freeze_panes != sheet_spec["freeze"]:
        rep.add("W-LAYOUT", "warn", ws.title, "-", f"freeze panes {ws.freeze_panes} != {sheet_spec['freeze']}")
    if bool(ws.sheet_view.showGridLines) != sheet_spec["show_grid_lines"]:
        rep.add("W-LAYOUT", "warn", ws.title, "-", "gridlines setting differs")
    actual = {get_column_letter(i): None for i in range(1, sheet_spec["ncols"] + 1)}
    for dim in ws.column_dimensions.values():
        for i in range(dim.min or 1, (dim.max or dim.min or 1) + 1):
            if i <= sheet_spec["ncols"] and dim.width:
                actual[get_column_letter(i)] = round(dim.width, 2)
    for letter, width in sheet_spec["widths"].items():
        if width and (actual.get(letter) is None or abs(actual[letter] - width) > 0.5):
            rep.add("W-LAYOUT", "warn", ws.title, f"{letter}:{letter}", f"column width {actual.get(letter)} != {width}")


def check_formulas(rep: Report, wb) -> None:
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and v.strip() in ERROR_VALUES:
                    rep.add("W-FORMULA", "error", ws.title, cell.coordinate, f"error value {v}")
                elif isinstance(v, str) and v.startswith("=") and any(e in v for e in ERROR_VALUES):
                    rep.add("W-FORMULA", "error", ws.title, cell.coordinate, f"formula contains error: {v}")


def check_passport_sheet(rep: Report, ws, spec_sp: dict, data: dict) -> None:
    sec = spec_sp["section"]
    ncols = spec_sp["ncols"]
    last_col = get_column_letter(ncols)
    title_rows = {s["title_row"] for s in data["sections"]}
    merged = {str(m) for m in ws.merged_cells.ranges}
    for s in data["sections"]:
        r = s["title_row"]
        if f"A{r}:{last_col}{r}" not in merged:
            rep.add("W-SECTION", "error", ws.title, f"A{r}", f"section title {s['agency']!r} not merged A:{last_col}")
        if s["header_row"] != r + 1:
            rep.add("W-SECTION", "error", ws.title, f"A{r + 1}", f"section {s['agency']!r}: header row must follow the title")
        else:
            _style_check(rep, ws, s["header_row"], sec["header_styles"], "error", "section header")
        _style_check(rep, ws, r, [sec["title_style"]], "error", "section title")
        for i, row in enumerate(s["rows"]):
            _style_check(rep, ws, row, sec["band_styles"][i % 2], "warn", f"body band {i % 2}")
        if s["rows"]:
            expected = list(range((s["header_row"] or r + 1) + 1, (s["header_row"] or r + 1) + 1 + len(s["rows"])))
            if s["rows"] != expected:
                rep.add("W-SECTION", "error", ws.title, f"A{s['rows'][0]}", f"section {s['agency']!r} has gaps inside its rows")
        else:
            rep.add("W-SECTION", "warn", ws.title, f"A{r}", f"section {s['agency']!r} has no rows")
    data_rows = {p["row"]: p for p in data["passports"]}
    for m in merged:
        start = m.split(":")[0]
        row = int(re.sub(r"[A-Z]+", "", start))
        if row in data_rows:
            p = data_rows[row]
            rep.add("W-MERGE", "error", ws.title, m,
                    f"merge hides passport values (row {row}: {p.get('service_name')!r}); unmerge the row")
        elif row not in title_rows:
            rep.add("W-MERGE", "error", ws.title, m, "merged range is not on a section title row")
    for p in data["passports"]:
        if p["section"] is None:
            rep.add("W-SECTION", "error", ws.title, f"A{p['row']}", "row outside any section")
        elif p.get("agency") != p["section"]:
            rep.add("W-CONSIST", "error", ws.title, f"B{p['row']}",
                    f"agency {p.get('agency')!r} in section {p['section']!r}")
        for f, header in c.PASSPORT_COLUMNS:
            v = p.get(f)
            col = get_column_letter(c.PASSPORT_FIELDS.index(f) + 1)
            if v is None or str(v).strip() == "":
                rep.add("W-EMPTY", "error", ws.title, f"{col}{p['row']}", f"{header} is empty")
            elif str(v).strip() in c.DEPRECATED_PLACEHOLDERS:
                rep.add("W-TERM", "error", ws.title, f"{col}{p['row']}",
                        f"{header} uses {str(v).strip()!r}; use {c.DEPRECATED_PLACEHOLDERS[str(v).strip()]!r}")
    keys = Counter((c.normalise_name(p.get("agency")), c.normalise_name(p.get("service_name"))) for p in data["passports"])
    for (agency, name), n in keys.items():
        if n > 1:
            rep.add("W-DUP", "error", ws.title, "-", f"service {name!r} appears {n}x in {agency!r}")
    ids = Counter(eservices_id_from_link(p.get("source_link")) for p in data["passports"])
    for sid, n in ids.items():
        if sid and n > 1:
            rep.add("W-DUP", "error", ws.title, "-", f"eServices service {sid} linked by {n} rows")


def check_totals(rep: Report, data: dict) -> None:
    actual = Counter(p["section"] for p in data["passports"])
    total = sum(actual.values())
    ov = data["overview"]
    declared_total = (ov.get("Service passports") or {}).get("value")
    if declared_total != total:
        rep.add("W-TOTAL", "error", "Overview", f"B{(ov.get('Service passports') or {}).get('row', '?')}",
                f"Service passports = {declared_total} but sheet has {total} rows")
    included = [a for a in data["agencies"] if (a.get("Status") or "Included") == "Included"]
    declared_entities = (ov.get("Entities included") or {}).get("value")
    if declared_entities != len(included):
        rep.add("W-TOTAL", "error", "Overview", f"B{(ov.get('Entities included') or {}).get('row', '?')}",
                f"Entities included = {declared_entities} but Agencies lists {len(included)} included")
    for r in data["overview_rows"]:
        if r["service_count"] != actual.get(r["agency"], 0):
            rep.add("W-TOTAL", "error", "Overview", f"C{r['row']}",
                    f"{r['agency']}: Service Count {r['service_count']} but {actual.get(r['agency'], 0)} rows")
    ov_names = [r["agency"] for r in data["overview_rows"]]
    ag_names = [a.get("Entity") for a in included]
    if ov_names != ag_names:
        rep.add("W-CONSIST", "error", "Overview", "A:A", f"Overview agencies {ov_names} != Agencies {ag_names}")
    for s in data["sections"]:
        if s["agency"] not in ag_names:
            rep.add("W-CONSIST", "error", "Service Passport", f"A{s['title_row']}",
                    f"section {s['agency']!r} not an included agency")
    ministry = data.get("ministry")
    for a in data["agencies"]:
        if a.get("Ministry") != ministry:
            rep.add("W-CONSIST", "error", "Agencies", f"B{a['row']}", f"ministry {a.get('Ministry')!r} != {ministry!r}")
        m = re.match(r"\s*(\d+)\b", str(a.get("Notes") or ""))
        if m and a.get("Entity") in actual and int(m.group(1)) != actual[a["Entity"]]:
            rep.add("W-TOTAL", "error", "Agencies", f"G{a['row']}",
                    f"Notes say {m.group(1)} services but sheet has {actual[a['Entity']]} rows")
    numbers = {a.get("Ministry No.") for a in data["agencies"]}
    if len(numbers) > 1:
        rep.add("W-CONSIST", "error", "Agencies", "A:A", f"several Ministry No. values {sorted(map(str, numbers))}")


def validate_workbook(path: Path, spec: dict) -> dict:
    rep = Report()
    wb = openpyxl.load_workbook(path)
    if wb.sheetnames != c.SHEET_NAMES:
        rep.add("W-SHEETS", "error", "-", "-", f"sheets {wb.sheetnames} != {c.SHEET_NAMES}")
    data = read_workbook(path)
    if "Overview" in wb.sheetnames:
        ws, so = wb["Overview"], spec["sheets"]["Overview"]
        check_layout(rep, ws, so)
        hr = so["table"]["header_row"]
        if [ws.cell(hr, i).value for i in range(1, 4)] != c.OVERVIEW_TABLE_HEADERS:
            rep.add("W-HEADER", "error", "Overview", f"A{hr}", "agency table header differs")
        else:
            _style_check(rep, ws, hr, so["table"]["header_styles"], "error", "table header")
        if ws.cell(so["title"]["row"], 1).value != c.WORKBOOK_TITLE:
            rep.add("W-HEADER", "error", "Overview", f"A{so['title']['row']}", "title text differs")
    if "Agencies" in wb.sheetnames:
        ws, sa = wb["Agencies"], spec["sheets"]["Agencies"]
        check_layout(rep, ws, sa)
        hr = sa["header"]["row"]
        if [ws.cell(hr, i).value for i in range(1, sa["ncols"] + 1)] != c.AGENCIES_HEADERS:
            rep.add("W-HEADER", "error", "Agencies", f"A{hr}", "header differs")
        else:
            _style_check(rep, ws, hr, sa["header"]["styles"], "error", "header")
        for i, a in enumerate(data["agencies"]):
            _style_check(rep, ws, a["row"], sa["body"]["band_styles"][i % 2], "warn", f"body band {i % 2}")
            if a.get("Status") not in c.AGENCY_STATUS:
                rep.add("W-TERM", "error", "Agencies", f"F{a['row']}", f"status {a.get('Status')!r}")
    if "Service Passport" in wb.sheetnames:
        ws, sp = wb["Service Passport"], spec["sheets"]["Service Passport"]
        check_layout(rep, ws, sp)
        check_passport_sheet(rep, ws, sp, data)
    check_totals(rep, data)
    check_formulas(rep, wb)
    errors = sum(1 for x in rep.items if x["severity"] == "error")
    by_rule = Counter(x["rule"] for x in rep.items)
    log.info("validated %s errors=%d warnings=%d", path, errors, len(rep.items) - errors)
    return {"ok": errors == 0, "workbook": str(path), "errors": errors, "warnings": len(rep.items) - errors,
            "by_rule": dict(by_rule), "rows": len(data["passports"]), "findings": rep.items}


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("workbook", type=Path)
    ap.add_argument("--spec", type=Path, default=templates_dir() / "workbook_spec.json")
    ap.add_argument("--out", type=Path)
    ap.add_argument("--summary", action="store_true", help="print only counts per rule, not every finding")
    args = ap.parse_args()
    try:
        report = validate_workbook(args.workbook, load_json(args.spec))
    except (OSError, KeyError, ValueError) as exc:
        log.error("cannot validate %s: %s", args.workbook, exc)
        return 2
    if args.out:
        dump_json(report, args.out)
    printed = {k: v for k, v in report.items() if k != "findings"} if args.summary else report
    print(json.dumps(printed, ensure_ascii=False, indent=2))
    return 0 if report["ok"] else 1


if __name__ == "__main__":
    sys.exit(main())
