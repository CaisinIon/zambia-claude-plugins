#!/usr/bin/env python3
"""Registry of services DotGov already built (the ZIGS service list behind Zambia eServices).

These services are not researched or verified: each one gets a placeholder passport whose
fields carry `FROM DOTGOV [<agency> - <service name>] {{DOTGOV:<ServiceID>:<field>}}`,
filled later from a database export (see dotgov_fill.py). ServiceID equals the eServices ID.

Commands:
  extract <xlsx> [--out FILE]                    build the registry JSON from the DotGov Excel
  show [--ministry NAME | --agency NAME]         ministries with their agencies and service counts; with
                                                 --ministry or --agency: the services by name
  match --agency NAME --name SERVICE [--run RUN --slug SLUG]
                                                 is a web-found service a DotGov service?

Registry lookup: <project>/input/dotgov_services.json (the visible copy, next to ministries.json),
else templates/dotgov_services.json (bundled with the skill; used by plugin installs).
"""
from __future__ import annotations

import argparse
import hashlib
import re
import sys
import unicodedata
from datetime import datetime, timezone
from pathlib import Path

from rapidfuzz import fuzz

import zmcontract as c
from zmlog import dump_json, get_logger, kv, load_json, project_root, templates_dir

log = get_logger("dotgov")

REGISTRY_FILE = "dotgov_services.json"
MINISTRIES_FILE = "dotgov_ministries.json"   # agency -> ministry map, applied by extract
SHEET_SERVICES = "Sheet1"
SHEET_QUERY = "Sheet2"
HEADERS = ["ServiceID", "name", "Description", "Department/Agency"]
TOKEN_RE = re.compile(r"\{\{DOTGOV:(\d+):([a-z_]+)\}\}")
ANY_TOKEN_RE = re.compile(r"\{\{\s*DOTGOV[^}]*\}\}", re.I)

_cache: dict[tuple[str, float], dict] = {}
_warned_missing = False


# ---------------------------------------------------------------- extract

def clean(text) -> str:
    """NFC, trimmed, internal whitespace runs collapsed to one space; curly quotes kept."""
    text = unicodedata.normalize("NFC", "" if text is None else str(text))
    return re.sub(r"\s+", " ", text).strip()


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def agency_ministries(path: Path | None = None) -> dict[str, str]:
    """{DotGov agency name: ministry} from templates/dotgov_ministries.json ({} when absent)."""
    path = Path(path) if path else templates_dir() / MINISTRIES_FILE
    if not path.is_file():
        log.warning("no agency->ministry map at %s; registry gets no ministry names", path)
        return {}
    return {a: v["ministry"] for a, v in load_json(path)["agencies"].items()}


def extract(xlsx: Path, ministries: Path | None = None) -> dict:
    """Read the DotGov Excel (Sheet1 services, Sheet2 SQL query) into the registry shape.
    Agencies and services get their ministry from the agency->ministry map; both lists are
    sorted ministry -> agency -> service name so the file reads by name."""
    from openpyxl import load_workbook

    xlsx = Path(xlsx)
    log.info("extract start %s", kv(file=str(xlsx)))
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    if SHEET_SERVICES not in wb.sheetnames:
        raise ValueError(f"{xlsx.name}: sheet {SHEET_SERVICES!r} not found (sheets: {wb.sheetnames})")
    rows = list(wb[SHEET_SERVICES].iter_rows(values_only=True))
    found = [clean(h) for h in (rows[0] if rows else [])][:len(HEADERS)]
    if found != HEADERS:
        raise ValueError(f"{xlsx.name}: expected headers {HEADERS}, found {found}")
    query = ""
    if SHEET_QUERY in wb.sheetnames:
        lines = [str(r[0]) for r in wb[SHEET_QUERY].iter_rows(values_only=True) if r and r[0] is not None]
        query = "\n".join(ln.rstrip() for ln in lines)

    services, seen = [], set()
    for n, row in enumerate(rows[1:], start=2):
        if not row or all(v is None or clean(v) == "" for v in row[:4]):
            log.debug("row %d empty; skipped", n)
            continue
        raw_id = row[0]
        try:
            sid = int(raw_id)
            if sid != float(raw_id):
                raise ValueError
        except (TypeError, ValueError):
            raise ValueError(f"{xlsx.name} row {n}: ServiceID {raw_id!r} is not an integer") from None
        if sid in seen:
            raise ValueError(f"{xlsx.name} row {n}: duplicate ServiceID {sid}")
        seen.add(sid)
        name, desc, agency = clean(row[1]), clean(row[2]), clean(row[3])
        if not name or not agency:
            raise ValueError(f"{xlsx.name} row {n}: ServiceID {sid} has no name or agency")
        services.append({
            "service_id": sid, "name": name, "name_key": c.normalise_name(name),
            "description": desc, "agency": agency, "agency_key": c.agency_key(agency),
            "ministry": None, "eservices_link": c.ESERVICES_SERVICE_URL.format(id=sid),
        })
        log.debug("row %d %s", n, kv(id=sid, name=name, agency=agency))
    owner = agency_ministries(ministries)
    unmapped = sorted({s["agency"] for s in services if s["agency"] not in owner})
    if owner and unmapped:
        log.warning("agencies without a ministry: %s", unmapped)
    for s in services:
        s["ministry"] = owner.get(s["agency"])
    services.sort(key=lambda s: ((s["ministry"] or "~"), s["agency"].lower(), s["name"].lower(), s["service_id"]))

    by_agency: dict[str, list[int]] = {}
    for s in services:
        by_agency.setdefault(s["agency"], []).append(s["service_id"])
    agencies = [{"name": a, "key": c.agency_key(a), "ministry": owner.get(a), "service_count": len(ids),
                 "service_ids": ids}
                for a, ids in sorted(by_agency.items(), key=lambda kv: (owner.get(kv[0]) or "~", kv[0].lower()))]
    sha = _sha256(xlsx)
    reg = {
        "schema_version": 1,
        "source": {"file": xlsx.name, "sheet": SHEET_SERVICES, "sha256": sha,
                   "extracted_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
                   "rows": len(services), "query": query},
        "placeholder": {"marker": c.DOTGOV_MARKER, "token": c.DOTGOV_TOKEN, "fields": list(c.DOTGOV_FIELDS)},
        "agencies": agencies,
        "services": services,
    }
    log.info("extract done %s", kv(rows=len(services), agencies=len(agencies),
                                   ministries=len({s["ministry"] for s in services} - {None}), sha256=sha))
    return reg


# ---------------------------------------------------------------- load / lookup

def registry_path() -> Path | None:
    """Project override first, then the copy bundled with the skill."""
    for path in (project_root() / "input" / REGISTRY_FILE, templates_dir() / REGISTRY_FILE):
        if path.is_file():
            return path
    return None


def load(path: Path | str | None = None) -> dict | None:
    """Load and validate the registry; None (with one WARN) when no registry exists."""
    global _warned_missing
    path = Path(path) if path else registry_path()
    if path is None or not path.is_file():
        if not _warned_missing:
            log.warning("registry not found; DotGov skip disabled %s", kv(path=str(path) if path else None))
            _warned_missing = True
        return None
    key = (str(path.resolve()), path.stat().st_mtime)
    if key in _cache:
        return _cache[key]
    from zmschema import schema_errors

    data = load_json(path)
    errors = schema_errors(data, "dotgov_services")
    if errors:
        raise ValueError(f"{path}: invalid DotGov registry: {errors[:5]}")
    data["_path"] = str(path)
    data["_by_id"] = {s["service_id"]: s for s in data["services"]}
    _cache[key] = data
    log.info("load %s", kv(source=str(path), services=len(data["services"]), sha256=data["source"]["sha256"]))
    return data


def for_run(meta: dict | None) -> dict | None:
    """The registry for a run, from its run.json (None when started with --no-dotgov).
    No run.json -> the default registry. A registry changed since run init is used, with a WARN."""
    conf = (meta or {}).get("dotgov") or {}
    if conf.get("enabled") is False:
        log.debug("DotGov disabled for this run")
        return None
    path = Path(conf["path"]) if conf.get("path") and Path(conf["path"]).is_file() else None
    reg = load(path)
    if reg and conf.get("sha256") and reg["source"]["sha256"] != conf["sha256"]:
        log.warning("DotGov registry changed since run init (run %s, now %s)",
                    conf["sha256"][:12], reg["source"]["sha256"][:12])
    return reg


def for_run_dir(run: Path | None) -> dict | None:
    """for_run() for a run folder; a folder without run.json gets the default registry."""
    meta = load_json(Path(run) / "run.json") if run and (Path(run) / "run.json").is_file() else None
    return for_run(meta)


def by_id(reg: dict | None) -> dict[int, dict]:
    if not reg:
        return {}
    return reg.get("_by_id") or {s["service_id"]: s for s in reg["services"]}


def marker(service: dict) -> str:
    return c.DOTGOV_MARKER.format(agency=service["agency"], service_name=service["name"])


def token(service_id: int, field: str) -> str:
    return c.DOTGOV_TOKEN.format(service_id=service_id, field=field)


def placeholder(service: dict, field: str) -> str:
    """Exact cell text for one DotGov placeholder field: marker + token."""
    return f"{marker(service)} {token(service['service_id'], field)}"


CELL_RE = re.compile(r"^FROM DOTGOV \[(.+?) - (.+)\] \{\{DOTGOV:(\d+):([a-z_]+)\}\}$")


def parse_cell(text) -> dict | None:
    """Split a placeholder cell back into {agency, service_name, service_id, field}; None if not one."""
    m = CELL_RE.match(str(text or "").strip())
    if not m:
        return None
    return {"agency": m.group(1), "service_name": m.group(2), "service_id": int(m.group(3)), "field": m.group(4)}


def is_placeholder(p: dict) -> bool:
    return bool(p.get("dotgov")) or p.get("verification") == c.DOTGOV_VERIFICATION


# ---------------------------------------------------------------- matching

def _agency_names(agency: dict | str) -> list[str]:
    if isinstance(agency, str):
        return [agency]
    names = [agency.get("official_name"), agency.get("roster_name"), *(agency.get("aliases") or [])]
    return [n for n in names if n]


def candidates(reg: dict, agency: dict | str) -> list[dict]:
    """DotGov services that may belong to this agency: its eServices IDs, plus agencies
    whose DotGov name fuzzy-matches the agency's name or aliases."""
    ids = set()
    if isinstance(agency, dict):
        ids |= set((agency.get("eservices") or {}).get("service_ids") or [])
        # services reassigned to this agency are not in its own service_ids
        ids |= {p["eservices_id"] for p in agency.get("passports") or [] if p.get("eservices_id")}
    keys = [c.agency_key(n) for n in _agency_names(agency)]
    out = []
    for s in reg["services"]:
        if s["service_id"] in ids or any(k and fuzz.token_set_ratio(k, s["agency_key"]) >= c.DOTGOV_MATCH
                                         for k in keys):
            out.append(s)
    log.debug("candidates %s", kv(agency=_agency_names(agency)[:1], by_id=len(ids), total=len(out)))
    return out


def score(a: str, b: str) -> float:
    """Mean of token-set and token-sort ratios: word order and spelling variants still score
    high, but a name that is only a subset ("Licence" vs "Casino Licence") drops below a match."""
    a, b = c.normalise_name(a), c.normalise_name(b)
    if not a or not b:
        return 0.0
    return (fuzz.token_set_ratio(a, b) + fuzz.token_sort_ratio(a, b)) / 2


def match(reg: dict, agency: dict | str, name: str) -> dict:
    """Best DotGov match for a service name within one agency: status match|possible|none."""
    best, best_score = None, 0.0
    for s in candidates(reg, agency):
        sc = score(name, s["name"])
        if sc > best_score:
            best, best_score = s, sc
    status = ("match" if best_score >= c.DOTGOV_MATCH else
              "possible" if best_score >= c.DOTGOV_POSSIBLE else "none")
    result = {"status": status, "service_id": best["service_id"] if best and status != "none" else None,
              "name": best["name"] if best and status != "none" else None,
              "agency": best["agency"] if best and status != "none" else None,
              "score": round(best_score, 1), "query": name}
    log.debug("match %s", kv(agency=_agency_names(agency)[:1], name=name, best=best and best["name"],
                             score=result["score"], status=status))
    return result


# ---------------------------------------------------------------- CLI

def _emit(data) -> None:
    import json
    json.dump(data, sys.stdout, ensure_ascii=False, indent=2)
    sys.stdout.write("\n")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    ep = sub.add_parser("extract")
    ep.add_argument("xlsx", type=Path)
    ep.add_argument("--out", type=Path, help=f"default: input/{REGISTRY_FILE} (visible copy) and templates/{REGISTRY_FILE}")
    sp = sub.add_parser("show")
    sp.add_argument("--agency")
    sp.add_argument("--ministry")
    sp.add_argument("--registry", type=Path)
    mp = sub.add_parser("match")
    mp.add_argument("--agency", required=True)
    mp.add_argument("--name", required=True)
    mp.add_argument("--run", type=Path)
    mp.add_argument("--slug")
    mp.add_argument("--registry", type=Path)
    args = ap.parse_args()

    try:
        if args.cmd == "extract":
            reg = extract(args.xlsx)
            outs = [args.out] if args.out else [project_root() / "input" / REGISTRY_FILE, templates_dir() / REGISTRY_FILE]
            for out in outs:
                dump_json(reg, out)
                log.info("wrote %s", out)
            _emit({"out": [str(o) for o in outs], "rows": reg["source"]["rows"], "agencies": len(reg["agencies"]),
                   "sha256": reg["source"]["sha256"]})
            return 0
        reg = load(args.registry)
        if reg is None:
            log.error("no DotGov registry; run: dotgov_registry.py extract <xlsx>")
            return 2
        if args.cmd == "show":
            if args.ministry:
                key = c.normalise_name(args.ministry)
                rows = [s for s in reg["services"] if c.normalise_name(s.get("ministry")) == key]
                _emit({"ministry": args.ministry, "services": len(rows),
                       "by_agency": {a: [s["name"] for s in rows if s["agency"] == a]
                                     for a in dict.fromkeys(s["agency"] for s in rows)}})
            elif args.agency:
                _emit([{k: s.get(k) for k in ("service_id", "name", "agency", "ministry")}
                       for s in sorted(candidates(reg, args.agency), key=lambda s: s["name"].lower())])
            else:
                tree: dict[str, dict[str, int]] = {}
                for a in reg["agencies"]:
                    tree.setdefault(a.get("ministry") or "(no ministry)", {})[a["name"]] = a["service_count"]
                _emit({"source": reg["source"]["file"], "sha256": reg["source"]["sha256"],
                       "services": len(reg["services"]),
                       "ministries": {m: {"services": sum(v.values()), "agencies": v} for m, v in tree.items()}})
        elif args.cmd == "match":
            agency: dict | str = args.agency
            if args.run and args.slug:
                agency = load_json(args.run / "agencies" / f"{args.slug}.json")
                agency.setdefault("aliases", []).append(args.agency)
            _emit(match(reg, agency, args.name))
    except ValueError as exc:
        log.error("%s", exc)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
