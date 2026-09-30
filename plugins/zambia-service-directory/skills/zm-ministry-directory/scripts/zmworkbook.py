"""Read and write ministry workbooks in the fixed reference format (workbook_spec.json)."""
from __future__ import annotations

import re
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

import zmcontract as c
from zmlog import get_logger
from zmstyle import apply_style, estimate_row_height

log = get_logger("workbook")
SECTION_PREFIX = c.SECTION_TITLE.split("{")[0]
ESERVICES_ID_RE = re.compile(r"eservices\.gov\.zm/#/service/(\d+)")


# ------------------------------------------------------------------ reading

class _Grid:
    """Cell values of one sheet loaded in read-only mode.

    Read-only mode keeps values hidden under merged ranges; a normal openpyxl
    load discards them (the reference workbook had 5 passports hidden this way).
    """

    def __init__(self, rows: list[tuple]):
        self.rows = rows
        self.max_row = len(rows)

    def cell(self, r: int, col: int):
        return _Value(self.rows[r - 1][col - 1] if r - 1 < len(self.rows) and col - 1 < len(self.rows[r - 1]) else None)


class _Value:
    __slots__ = ("value",)

    def __init__(self, value):
        self.value = value


def _grid(path: Path) -> dict[str, _Grid]:
    wb = openpyxl.load_workbook(path, read_only=True)
    grids = {ws.title: _Grid([tuple(r) for r in ws.iter_rows(values_only=True)]) for ws in wb.worksheets}
    wb.close()
    return grids


def _row_values(ws, r: int, n: int) -> list:
    return [ws.cell(r, col).value for col in range(1, n + 1)]


def eservices_id_from_link(link) -> int | None:
    m = ESERVICES_ID_RE.search(str(link or ""))
    return int(m.group(1)) if m else None


def read_workbook(path: Path) -> dict:
    """Parse a ministry workbook into plain data (tolerant of layout drift).

    Returns {ministry, scope, overview: {...}, overview_rows, agencies, sections, passports}.
    passports: dicts keyed by PASSPORT_FIELDS plus 'row' and 'section'.
    """
    wb = _grid(path)
    out: dict = {"path": str(path), "sheet_names": list(wb), "overview": {}, "overview_rows": [],
                 "agencies": [], "sections": [], "passports": []}
    if "Overview" in wb:
        ws = wb["Overview"]
        header_row = None
        for r in range(1, ws.max_row + 1):
            label = ws.cell(r, 1).value
            if label in c.OVERVIEW_KEYS:
                out["overview"][label] = {"value": ws.cell(r, 2).value, "row": r}
            if _row_values(ws, r, 3) == c.OVERVIEW_TABLE_HEADERS:
                header_row = r
        if header_row:
            for r in range(header_row + 1, ws.max_row + 1):
                vals = _row_values(ws, r, 3)
                if vals[0]:
                    out["overview_rows"].append({"agency": vals[0], "entity_type": vals[1], "service_count": vals[2], "row": r})
    out["ministry"] = (out["overview"].get("Ministry") or {}).get("value")
    out["scope"] = (out["overview"].get("Scope") or {}).get("value")

    if "Agencies" in wb:
        ws = wb["Agencies"]
        n = len(c.AGENCIES_HEADERS)
        header_row = next((r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == c.AGENCIES_HEADERS[0]), None)
        if header_row:
            for r in range(header_row + 1, ws.max_row + 1):
                vals = _row_values(ws, r, n)
                if any(v is not None for v in vals):
                    out["agencies"].append({**dict(zip(c.AGENCIES_HEADERS, vals)), "row": r})

    if "Service Passport" in wb:
        ws = wb["Service Passport"]
        n = len(c.PASSPORT_HEADERS)
        section = None
        for r in range(1, ws.max_row + 1):
            first = ws.cell(r, 1).value
            vals = _row_values(ws, r, n)
            if isinstance(first, str) and first.startswith(SECTION_PREFIX):
                section = {"agency": first[len(SECTION_PREFIX):].strip(), "title_row": r, "header_row": None,
                           "rows": []}
                out["sections"].append(section)
            elif vals == c.PASSPORT_HEADERS:
                if section is not None:
                    section["header_row"] = r
            elif any(v is not None for v in vals):
                row = dict(zip(c.PASSPORT_FIELDS, vals))
                row.update({"row": r, "section": section["agency"] if section else None})
                out["passports"].append(row)
                if section is not None:
                    section["rows"].append(r)
    log.debug("read %s: agencies=%d sections=%d passports=%d", path, len(out["agencies"]),
              len(out["sections"]), len(out["passports"]))
    return out


# ------------------------------------------------------------------ writing

def _set_widths(ws, widths: dict) -> None:
    for letter, width in widths.items():
        if width:
            ws.column_dimensions[letter].width = width


def _write_row(ws, r: int, values: list, styles: list[dict]) -> None:
    for col, (value, style) in enumerate(zip(values, styles), start=1):
        cell = ws.cell(r, col, value)
        apply_style(cell, style)


def _title(ws, r: int, text: str, spec_title: dict, ncols: int) -> None:
    cell = ws.cell(r, 1, text)
    apply_style(cell, spec_title["style"] if "style" in spec_title else spec_title)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)


def _band(styles_pair: list, index: int) -> list[dict]:
    return styles_pair[index % 2]


def write_workbook(path: Path, spec: dict, data: dict) -> None:
    """data: {ministry, ministry_no, scope, agencies: [agency rows], sections: [{agency, rows: [passport dicts]}]}.

    agency row keys: official_name, entity_type, digitizable_service_areas, status, notes, source_link.
    """
    wb = openpyxl.Workbook()
    try:
        df = spec["default_font"]
        wb._named_styles["Normal"].font = Font(name=df["name"], sz=df["size"])
    except Exception:  # noqa: BLE001 - cosmetic only
        log.debug("could not set Normal font")
    wb.remove(wb.active)
    ministry = data["ministry"]

    # Overview
    so = spec["sheets"]["Overview"]
    ws = wb.create_sheet("Overview")
    ws.sheet_view.showGridLines = so["show_grid_lines"]
    _set_widths(ws, so["widths"])
    tr = so["title"]["row"]
    _title(ws, tr, c.WORKBOOK_TITLE, so["title"], so["title"]["merge_cols"])
    ws.row_dimensions[tr].height = so["title"]["height"]
    included = [a for a in data["agencies"] if a["status"] == "Included"]
    total = sum(len(s["rows"]) for s in data["sections"])
    key_values = {"Ministry": ministry, "Entities included": len(included), "Service passports": total,
                  "Scope": data.get("scope") or c.DEFAULT_SCOPE}
    keys = so["keys"]
    for i, label in enumerate(keys["labels"]):
        r = keys["first_row"] + i
        apply_style(ws.cell(r, 1, label), keys["label_style"])
        apply_style(ws.cell(r, 2, key_values[label]), keys["value_style"])
        if keys["heights"].get(label):
            ws.row_dimensions[r].height = keys["heights"][label]
    table = so["table"]
    hr = table["header_row"]
    _write_row(ws, hr, table["headers"], table["header_styles"])
    ws.row_dimensions[hr].height = table["header_height"]
    counts = {s["agency"]: len(s["rows"]) for s in data["sections"]}
    for i, a in enumerate(included):
        r = hr + 1 + i
        _write_row(ws, r, [a["official_name"], a["entity_type"], counts.get(a["official_name"], 0)],
                   _band(table["band_styles"], i))
        ws.row_dimensions[r].height = table["body_height"]
    ws.freeze_panes = so["freeze"]

    # Agencies
    sa = spec["sheets"]["Agencies"]
    ws = wb.create_sheet("Agencies")
    ws.sheet_view.showGridLines = sa["show_grid_lines"]
    _set_widths(ws, sa["widths"])
    widths = [sa["widths"].get(get_column_letter(i)) for i in range(1, sa["ncols"] + 1)]
    tr = sa["title"]["row"]
    _title(ws, tr, c.AGENCIES_TITLE.format(ministry=ministry), sa["title"], sa["ncols"])
    ws.row_dimensions[tr].height = sa["title"]["height"]
    hr = sa["header"]["row"]
    _write_row(ws, hr, sa["header"]["headers"], sa["header"]["styles"])
    ws.row_dimensions[hr].height = sa["header"]["height"]
    for i, a in enumerate(data["agencies"]):
        r = sa["body"]["first_row"] + i
        values = [data["ministry_no"], ministry, a["official_name"], a["entity_type"],
                  a["digitizable_service_areas"], a["status"], a["notes"], a["source_link"]]
        _write_row(ws, r, values, _band(sa["body"]["band_styles"], i))
        ws.row_dimensions[r].height = estimate_row_height(values, widths, sa["body"]["min_height"])
    ws.freeze_panes = sa["freeze"]

    # Service Passport
    sp = spec["sheets"]["Service Passport"]
    sec = sp["section"]
    ws = wb.create_sheet("Service Passport")
    ws.sheet_view.showGridLines = sp["show_grid_lines"]
    _set_widths(ws, sp["widths"])
    widths = [sp["widths"].get(get_column_letter(i)) for i in range(1, sp["ncols"] + 1)]
    r = sp["first_row"]
    for s_index, section in enumerate(s for s in data["sections"] if s["rows"]):
        if s_index:
            ws.row_dimensions[r].height = sec["spacer_height"]
            r += 1
        cell = ws.cell(r, 1, sec["title_text"].format(agency=section["agency"]))
        apply_style(cell, sec["title_style"])
        for col in range(2, sp["ncols"] + 1):
            apply_style(ws.cell(r, col), sec["title_style"])
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=sp["ncols"])
        ws.row_dimensions[r].height = sec["title_height"]
        r += 1
        _write_row(ws, r, sec["headers"], sec["header_styles"])
        ws.row_dimensions[r].height = sec["header_height"]
        r += 1
        for i, p in enumerate(section["rows"]):
            values = [p.get(f) for f in c.PASSPORT_FIELDS]
            _write_row(ws, r, values, _band(sec["band_styles"], i))
            ws.row_dimensions[r].height = estimate_row_height(values, widths, sec["min_body_height"])
            r += 1
    ws.freeze_panes = sp["freeze"]

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    log.info("workbook saved %s agencies=%d passports=%d", path, len(data["agencies"]), total)
