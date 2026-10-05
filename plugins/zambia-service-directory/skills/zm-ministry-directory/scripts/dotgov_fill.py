#!/usr/bin/env python3
"""Fill DotGov placeholders in a ministry workbook from a DotGov database export.

Usage:
  dotgov_fill.py fill <workbook.xlsx> --data <export.json|export.csv> [--out FILE] [--report FILE]

Every Service Passport cell holding `FROM DOTGOV [...] {{DOTGOV:<ServiceID>:<field>}}` is replaced, as a
whole cell, by the export's value for (ServiceID, field). A missing or empty value keeps the cell and is
listed under `missing`. Styles are kept (values only are changed).

Export: one row per service, keyed by `ServiceID` (or `service_id`). Field columns use the passport field
keys (who_can_apply, eligibility_requirements, fee, processing_time, validity, legal_references) or the
workbook headers (Who Can Apply, Fee, ...). CSV must be UTF-8; JSON is a list of objects or {"services": [...]}.

Default --out: <workbook>_filled.xlsx next to the input. The input is overwritten only when --out names it.
The result is checked with validate_workbook.py; exit 0 = filled and valid, 1 = workbook errors, 2 = bad input.
"""
from __future__ import annotations

import argparse
import csv
import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import MergedCell

import dotgov_registry
import zmcontract as c
from zmlog import dump_json, get_logger, kv, load_json, templates_dir

log = get_logger("dotgov_fill")
ID_KEYS = ("ServiceID", "service_id", "serviceid", "ID")
HEADER_TO_FIELD = {h.lower(): f for f, h in c.PASSPORT_COLUMNS}


def _field_key(key: str) -> str | None:
    k = str(key or "").strip()
    if k in c.DOTGOV_FIELDS:
        return k
    f = HEADER_TO_FIELD.get(k.lower())
    return f if f in c.DOTGOV_FIELDS else None


def load_export(path: Path) -> dict[int, dict[str, str]]:
    """{ServiceID: {field: value}} from a CSV or JSON export. Unknown columns are ignored."""
    path = Path(path)
    if path.suffix.lower() == ".csv":
        with open(path, encoding="utf-8-sig", newline="") as fh:
            rows = list(csv.DictReader(fh))
    elif path.suffix.lower() == ".json":
        data = load_json(path)
        rows = data.get("services", []) if isinstance(data, dict) else data
    else:
        raise ValueError(f"{path.name}: export must be .csv or .json")
    out: dict[int, dict[str, str]] = {}
    for n, row in enumerate(rows, start=1):
        raw = next((row[k] for k in ID_KEYS if k in row and row[k] not in (None, "")), None)
        try:
            sid = int(float(raw))
        except (TypeError, ValueError):
            raise ValueError(f"{path.name} row {n}: no integer ServiceID ({raw!r})") from None
        values = {}
        for k, v in row.items():
            f = _field_key(k)
            if f and v is not None and str(v).strip():
                values[f] = str(v).strip()
        if sid in out:
            log.warning("export row %d repeats ServiceID %d; later values win", n, sid)
        out.setdefault(sid, {}).update(values)
    log.info("export loaded %s", kv(file=str(path), services=len(out)))
    return out


def fill(workbook: Path, export: dict[int, dict[str, str]], out: Path) -> dict:
    wb = openpyxl.load_workbook(workbook)
    if "Service Passport" not in wb.sheetnames:
        raise ValueError(f"{workbook.name}: no 'Service Passport' sheet")
    ws = wb["Service Passport"]
    filled, missing, skipped, seen = 0, [], [], set()
    for row in ws.iter_rows():
        for cell in row:
            if not isinstance(cell.value, str) or "DOTGOV" not in cell.value:
                continue
            parsed = dotgov_registry.parse_cell(cell.value)
            if not parsed:
                log.warning("%s: not a well-formed DotGov placeholder; left as is: %r", cell.coordinate, cell.value)
                skipped.append({"cell": cell.coordinate, "value": cell.value})
                continue
            sid, field = parsed["service_id"], parsed["field"]
            seen.add(sid)
            value = (export.get(sid) or {}).get(field)
            if not value:
                missing.append({"cell": cell.coordinate, "service_id": sid, "field": field})
                log.debug("%s %s:%s no value in export", cell.coordinate, sid, field)
                continue
            if isinstance(cell, MergedCell):
                skipped.append({"cell": cell.coordinate, "value": cell.value, "reason": "merged cell"})
                continue
            cell.value = value
            filled += 1
            log.debug("%s %s:%s filled", cell.coordinate, sid, field)
    unknown = sorted(set(export) - seen)
    if unknown:
        log.warning("export ServiceIDs not in this workbook: %s", unknown[:20])
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    result = {"workbook": str(workbook), "out": str(out), "filled": filled, "missing": missing,
              "missing_count": len(missing), "unknown_ids": unknown, "skipped": skipped,
              "services_in_workbook": len(seen)}
    log.info("fill %s", kv(services=len(seen), filled=filled, missing=len(missing), unknown_ids=len(unknown),
                            skipped=len(skipped), out=str(out)))
    return result


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    fp = sub.add_parser("fill")
    fp.add_argument("workbook", type=Path)
    fp.add_argument("--data", type=Path, required=True)
    fp.add_argument("--out", type=Path)
    fp.add_argument("--report", type=Path)
    args = ap.parse_args()
    from validate_workbook import validate_workbook

    out = args.out or args.workbook.with_name(f"{args.workbook.stem}_filled.xlsx")
    try:
        result = fill(args.workbook, load_export(args.data), out)
    except (OSError, ValueError, json.JSONDecodeError) as exc:
        log.error("%s", exc)
        return 2
    check = validate_workbook(out, load_json(templates_dir() / "workbook_spec.json"))
    result["validation"] = {k: v for k, v in check.items() if k != "findings"}
    if args.report:
        dump_json(result, args.report)
    print(json.dumps({k: v for k, v in result.items() if k not in ("missing", "skipped")}, ensure_ascii=False, indent=2))
    return 0 if check["ok"] else 1


if __name__ == "__main__":
    sys.exit(main())
