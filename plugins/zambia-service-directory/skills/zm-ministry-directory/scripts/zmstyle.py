"""Serialise openpyxl cell styles to plain dicts and back, and compare them."""
from __future__ import annotations

import math

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def _rgb(color) -> str | None:
    if color is None:
        return None
    rgb = getattr(color, "rgb", None)
    return rgb if isinstance(rgb, str) else None


def _side(side: Side) -> dict | None:
    if side is None or side.style is None:
        return None
    return {"style": side.style, "color": _rgb(side.color)}


def style_of(cell) -> dict:
    f, a, b, fl = cell.font, cell.alignment, cell.border, cell.fill
    return {
        "font": {"name": f.name, "size": float(f.sz) if f.sz else None, "bold": bool(f.b),
                 "italic": bool(f.i), "underline": f.u, "color": _rgb(f.color)},
        "fill": {"type": fl.fill_type, "color": _rgb(fl.fgColor) if fl.fill_type else None},
        "border": {k: _side(getattr(b, k)) for k in ("left", "right", "top", "bottom")},
        "alignment": {"horizontal": a.horizontal, "vertical": a.vertical, "wrap": bool(a.wrap_text)},
        "number_format": cell.number_format,
    }


def apply_style(cell, style: dict) -> None:
    f = style["font"]
    cell.font = Font(name=f["name"], sz=f["size"], b=f["bold"], i=f["italic"], u=f["underline"], color=f["color"])
    fl = style["fill"]
    cell.fill = PatternFill(fill_type=fl["type"], fgColor=fl["color"], bgColor=fl["color"]) if fl["type"] else PatternFill()
    sides = {}
    for k, v in style["border"].items():
        sides[k] = Side(style=v["style"], color=v["color"]) if v else Side()
    cell.border = Border(**sides)
    al = style["alignment"]
    cell.alignment = Alignment(horizontal=al["horizontal"], vertical=al["vertical"], wrap_text=al["wrap"] or None)
    cell.number_format = style["number_format"]


def style_diff(actual: dict, expected: dict, prefix: str = "") -> list[str]:
    """List 'path: actual != expected' differences between two style dicts."""
    diffs = []
    for key, exp in expected.items():
        act = actual.get(key) if isinstance(actual, dict) else None
        path = f"{prefix}.{key}" if prefix else key
        if isinstance(exp, dict) and isinstance(act, dict):
            diffs.extend(style_diff(act, exp, path))
        elif isinstance(exp, float) or isinstance(act, float):
            if exp is None or act is None or abs(float(exp) - float(act)) > 0.01:
                if exp != act:
                    diffs.append(f"{path}: {act!r} != {exp!r}")
        elif act != exp:
            diffs.append(f"{path}: {act!r} != {exp!r}")
    return diffs


LINE_HEIGHT_PT = 10.8  # one wrapped line of Arial 9 in the reference workbook
MAX_ROW_HEIGHT = 409.0


def estimate_row_height(values: list, widths: list[float | None], min_height: float,
                        font_size: float = 9.0) -> float:
    """Row height that fits wrapped text: lines x 10.8 pt (scaled by font size)."""
    lines = 1
    for value, width in zip(values, widths):
        if value is None:
            continue
        chars_per_line = max(4.0, (width or 8.43) * 1.15 * (9.0 / font_size))
        text_lines = 0
        for para in str(value).split("\n"):
            text_lines += max(1, math.ceil(len(para) / chars_per_line))
        lines = max(lines, text_lines)
    height = LINE_HEIGHT_PT * (font_size / 9.0) * lines
    return round(min(MAX_ROW_HEIGHT, max(min_height, height)), 2)
